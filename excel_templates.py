"""Downloadable Excel templates mapped to the data model.

Each template has a README sheet (plain-language instructions) and a Data sheet
with headers, an example row, and dropdown validation where it helps.
"""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
WRAP = Alignment(wrap_text=True, vertical="top")


def _readme_sheet(wb, title, intro_lines, column_help):
    ws = wb.active
    ws.title = "README - start here"
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    row = 3
    for line in intro_lines:
        ws.cell(row=row, column=1, value=line).alignment = WRAP
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Column").font = Font(bold=True)
    ws.cell(row=row, column=2, value="What to enter").font = Font(bold=True)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    for col, help_text in column_help:
        row += 1
        ws.cell(row=row, column=1, value=col).font = Font(bold=True)
        ws.cell(row=row, column=2, value=help_text).alignment = WRAP
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 90


def _data_sheet(wb, headers, example_row, dropdowns=None):
    ws = wb.create_sheet("Data - fill this in")
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
    for i, v in enumerate(example_row, start=1):
        ws.cell(row=2, column=i, value=v)
    ws["A2"].font = Font(italic=True, color="808080")
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 30
    ws.sheet_properties.tabColor = "1F4E79"
    if dropdowns:
        for col_idx, options in dropdowns.items():
            col = get_column_letter(col_idx)
            dv = DataValidation(type="list", formula1='"%s"' % ",".join(options),
                                allow_blank=True)
            dv.error = "Please pick one of: %s" % ", ".join(options)
            dv.prompt = "Pick from the list"
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}500")
    # note row 2 is the example; tell users via the README
    return ws


def _workbook_bytes(wb):
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def positions_template():
    wb = Workbook()
    _readme_sheet(
        wb, "Three-way positions - how to fill this in",
        [
            "This sheet loads one utility's as-filed, Staff, and final (or stipulation) "
            "numbers for a rate case into the app.",
            "1. Fill in one row per stage. A full case has three rows: as_filed, staff, final_or_stip.",
            "2. Row 2 is an EXAMPLE - overwrite it or delete it before uploading.",
            "3. The Docket must already exist in the app (add it with the Case header template first).",
            "4. Money is in millions of dollars: write 48.8 for $48.8 million. You can also write $48.8M - both work.",
            "5. ROE is a percent: write 9.84 for 9.84%.",
            "6. Leave a cell blank if you don't know the number - blank is always OK except Docket and Stage.",
            "7. The last five columns (equity ratio through taxes) are the full revenue-requirement build. "
            "Fill them in and the scenario workbench can run one-click what-ifs straight off the as-filed numbers.",
        ],
        [
            ("Docket", "The case number, e.g. 24-0468-EL-AIR. Must match a case already in the app."),
            ("Stage", "One of: as_filed (what the utility asked for), staff (what PUCO Staff recommended), "
                      "final_or_stip (what the Commission approved, or the stipulation if no order yet)."),
            ("Period", "Optional. For multi-year plans, e.g. TY1 2027. Leave blank for single-year cases."),
            ("Revenue requirement ($M)", "Optional. Total revenue requirement in millions, e.g. 955.1."),
            ("Revenue change ($M)", "Optional. Annual change vs. current rates, e.g. 48.8. Use a minus sign for a decrease, e.g. -19."),
            ("ROE (%)", "Optional. Return on equity as a percent, e.g. 9.84."),
            ("Capital structure", "Optional free text, e.g. 50.5% equity / 49.5% debt."),
            ("Rate base ($M)", "Optional. Rate base in millions, e.g. 2037.9."),
            ("Notes", "Optional. Anything worth remembering, e.g. 'Settlement adopted' or 'midpoint of Staff range'."),
            ("Equity ratio (%)", "Optional. Percent of the capital structure that is equity, e.g. 50.88. "
                                 "Powers the scenario workbench's one-click what-if math."),
            ("Cost of debt (%)", "Optional. Pre-tax cost of debt as a percent, e.g. 5.25. "
                                 "Powers the scenario workbench."),
            ("Operating expenses ($M)", "Optional. Annual operating expenses in millions. "
                                        "Powers the scenario workbench."),
            ("Depreciation ($M)", "Optional. Annual depreciation expense in millions. "
                                  "Powers the scenario workbench."),
            ("Taxes ($M)", "Optional. Income plus other taxes in millions. "
                            "Powers the scenario workbench."),
        ],
    )
    headers = ["Docket", "Stage (as_filed / staff / final_or_stip)", "Period (optional)",
               "Revenue requirement ($M, optional)", "Revenue change ($M, optional)",
               "ROE (%, optional)", "Capital structure (optional)", "Rate base ($M, optional)",
               "Notes (optional)", "Equity ratio (%, optional)", "Cost of debt (%, optional)",
               "Operating expenses ($M, optional)", "Depreciation ($M, optional)",
               "Taxes ($M, optional)"]
    _data_sheet(wb, headers,
                ["24-0468-EL-AIR  <-- EXAMPLE: overwrite or delete this row",
                 "as_filed", "", "955.1", "42.3", "10.15", "54.43% equity / 45.57% debt",
                 "3088.4", "EXAMPLE ROW - delete before uploading",
                 "54.43", "5.25", "", "", ""],
                dropdowns={2: ["as_filed", "staff", "final_or_stip"]})
    return _workbook_bytes(wb)


def riders_template():
    wb = Workbook()
    _readme_sheet(
        wb, "Tariff and rider inventory - how to fill this in",
        [
            "This sheet loads a utility's tariffs and riders into the app, including what "
            "happens to each one in the next rate case.",
            "1. One row per tariff or rider.",
            "2. Row 2 is an EXAMPLE - overwrite it or delete it before uploading.",
            "3. The Utility must already exist in the app (AEP Ohio, Duke Energy Ohio, AES Ohio, FirstEnergy).",
            "4. 'What happens next' is the important one: does this rider roll into base rates "
            "at the next case, stay as its own separate rider, or is it undecided?",
            "5. Leave a cell blank if you don't know - blank is always OK except Utility and Name.",
        ],
        [
            ("Utility", "The utility name, e.g. AES Ohio. Must match a utility already in the app."),
            ("Name", "Rider or tariff name, e.g. Rider DCR."),
            ("Type", "tariff or rider."),
            ("Status", "effective (in rates now), proposed (asked for but not approved), or terminated (gone / set to zero)."),
            ("Annual revenue ($M)", "Optional. Rough annual dollars flowing through it, in millions."),
            ("What happens next", "roll_into_base (merges into base rates next case), standalone (stays its own rider), "
                                 "or undecided."),
            ("Notes", "Optional. E.g. 'Reset to zero 1/1/2027'."),
        ],
    )
    headers = ["Utility", "Name", "Type (tariff / rider)", "Status (effective / proposed / terminated)",
               "Annual revenue ($M, optional)", "What happens next (roll_into_base / standalone / undecided)",
               "Notes (optional)"]
    _data_sheet(wb, headers,
                ["AES Ohio  <-- EXAMPLE: overwrite or delete this row", "Rider DCR", "rider",
                 "proposed", "", "roll_into_base", "EXAMPLE ROW - delete before uploading"],
                dropdowns={3: ["tariff", "rider"],
                           4: ["effective", "proposed", "terminated"],
                           6: ["roll_into_base", "standalone", "undecided"]})
    return _workbook_bytes(wb)


def cases_template():
    wb = Workbook()
    _readme_sheet(
        wb, "Case header info - how to fill this in",
        [
            "This sheet adds new rate cases to the app (or updates cases already there).",
            "1. One row per case. The Docket is the ID - uploading the same docket twice just updates it.",
            "2. Row 2 is an EXAMPLE - overwrite it or delete it before uploading.",
            "3. Dates are YYYY-MM-DD, e.g. 2026-04-01. Leave blank if unknown.",
            "4. Case type: historical (traditional test-year rate case), prospective (HB 15 three-year / forecasted test period plan),",
            "   tariff (tariff amendment, e.g. ATA data-center tariffs), rider (rider proceedings, e.g. RDR),",
            "   or other (any other PUCO proceeding: AAM, GCR, UNC, etc.).",
            "5. Status: decided (final order issued), stipulated (settlement filed, no order yet), or pending.",
        ],
        [
            ("Docket", "The case number, e.g. 26-0347-EL-AIR. This is how the app matches rows."),
            ("Utility", "The utility name, e.g. FirstEnergy. Must match a utility already in the app."),
            ("Case type", "historical or prospective for rate cases; tariff, rider, or other for non-rate proceedings."),
            ("Status", "decided, stipulated, or pending."),
            ("Date filed", "Optional. YYYY-MM-DD."),
            ("Staff report date", "Optional. YYYY-MM-DD."),
            ("Order date", "Optional. YYYY-MM-DD."),
            ("Test year", "Optional free text, e.g. 12 months ending 5/31/2025."),
            ("Date certain", "Optional. YYYY-MM-DD."),
            ("Notes", "Optional."),
        ],
    )
    headers = ["Docket", "Utility", "Case type (historical / prospective / tariff / rider / other)",
               "Status (decided / stipulated / pending)", "Date filed (YYYY-MM-DD, optional)",
               "Staff report date (optional)", "Order date (optional)", "Test year (optional)",
               "Date certain (optional)", "Notes (optional)"]
    _data_sheet(wb, headers,
                ["26-0347-EL-AIR  <-- EXAMPLE: overwrite or delete this row", "FirstEnergy",
                 "prospective", "pending", "2026-05-22", "", "", "TY1 7/1/2027-6/30/2028",
                 "", "EXAMPLE ROW - delete before uploading"],
                dropdowns={3: ["historical", "prospective", "tariff", "rider", "other"],
                           4: ["decided", "stipulated", "pending"]})
    return _workbook_bytes(wb)


def earned_roe_template():
    wb = Workbook()
    _readme_sheet(
        wb, "Earned ROE by year - how to fill this in",
        [
            "This sheet loads a utility's ACTUAL earned ROE per calendar year - what it really "
            "earned, not what the Commission authorized. The app compares it against the "
            "authorized ROE from the rate case on the utility's main page.",
            "1. One row per utility per year. Uploading the same utility + year twice just updates it.",
            "2. Row 2 is an EXAMPLE - overwrite it or delete it before uploading.",
            "3. The Utility must already exist in the app.",
            "4. Earned ROE comes from annual earnings filings: PUCO annual earnings reports, "
            "the utility's 10-K, or SEET filings - not from the rate case docket.",
            "5. Enter ROE as a percent, e.g. 10.42 means 10.42%.",
        ],
        [
            ("Utility", "The utility name, e.g. AES Ohio. Must match a utility already in the app."),
            ("Year", "Calendar year, e.g. 2024."),
            ("Earned ROE (%)", "Actual earned return on equity for that year, as a percent (10.42 = 10.42%)."),
            ("Source", "Where the number came from, e.g. '2024 PUCO earnings report' or '2024 10-K'."),
            ("Notes", "Optional."),
        ],
    )
    headers = ["Utility", "Year", "Earned ROE (%)", "Source", "Notes (optional)"]
    _data_sheet(wb, headers,
                ["AES Ohio  <-- EXAMPLE: overwrite or delete this row", 2024, 10.42,
                 "2024 PUCO earnings report", "EXAMPLE ROW - delete before uploading"])
    return _workbook_bytes(wb)


def discovery_template():
    wb = Workbook()
    _readme_sheet(
        wb, "Discovery log - how to fill this in",
        [
            "This sheet loads a case's discovery items - every DR, interrogatory, document "
            "request, deposition, and motion - into the command center so Regulatory Affairs "
            "can assign them out, track them, and assemble responses on time.",
            "1. One row per discovery item. Uploading the same item twice just updates it "
            "(the app matches on Docket + Type + Requesting party + Set + Number).",
            "2. Row 2 is an EXAMPLE - overwrite it or delete it before uploading.",
            "3. The Docket must already exist in the app (add it with the Case header template first).",
            "4. Dates are YYYY-MM-DD, e.g. 2026-10-15.",
            "5. Assigned to is the person who owns the response - nobody owns it, nobody does it.",
        ],
        [
            ("Docket", "The case number, e.g. 26-0347-EL-AIR. Must match a case already in the app."),
            ("Type", "data_request (a DR), interrogatory, production (request for production), "
                     "admission, deposition, motion, or other."),
            ("Requesting party", "Who served it, e.g. Staff, OCC, OEG, Kroger."),
            ("Set", "Optional. The discovery set, e.g. Set 3."),
            ("Number", "Optional. The item number, e.g. 3-12."),
            ("Served (date)", "Optional. When it was served, YYYY-MM-DD."),
            ("Due (date)", "When the response is due, YYYY-MM-DD. This drives the bottleneck view."),
            ("Assigned to", "Who owns the response, e.g. Jane Smith."),
            ("Status", "open, in_progress, in_review, complete, or withdrawn."),
            ("Response summary", "Optional. Short note on where the response stands."),
            ("Notes", "Optional."),
        ],
    )
    headers = ["Docket", "Type", "Requesting party", "Set (optional)", "Number (optional)",
               "Served - YYYY-MM-DD (optional)", "Due - YYYY-MM-DD", "Assigned to",
               "Status (open / in_progress / in_review / complete / withdrawn)",
               "Response summary (optional)", "Notes (optional)"]
    _data_sheet(wb, headers,
                ["26-0347-EL-AIR  <-- EXAMPLE: overwrite or delete this row", "data_request",
                 "Staff", "Set 1", "1-04", "2026-09-15", "2026-10-06", "Jane Smith",
                 "in_progress", "Draft with Regulatory", "EXAMPLE ROW - delete before uploading"],
                dropdowns={2: ["data_request", "interrogatory", "production", "admission",
                                "deposition", "motion", "other"],
                           9: ["open", "in_progress", "in_review", "complete", "withdrawn"]})
    return _workbook_bytes(wb)


def experts_template():
    wb = Workbook()
    _readme_sheet(
        wb, "Expert witness roster - how to fill this in",
        [
            "This sheet loads a case's expert witnesses into the command center so testimony "
            "deadlines are tracked next to the discovery they depend on.",
            "1. One row per expert per case. Uploading the same name twice just updates it.",
            "2. Row 2 is an EXAMPLE - overwrite it or delete it before uploading.",
            "3. The Docket must already exist in the app.",
            "4. Testimony due is YYYY-MM-DD - this drives the bottleneck view.",
        ],
        [
            ("Docket", "The case number, e.g. 26-0347-EL-AIR. Must match a case already in the app."),
            ("Name", "The expert's name, e.g. Dr. Jane Smith."),
            ("Firm", "Optional. Their firm or affiliation."),
            ("Topic", "Optional. What they testify on, e.g. Cost of capital."),
            ("Side", "utility (our witness), staff, or intervenor."),
            ("Testimony due", "YYYY-MM-DD."),
            ("Status", "retained, drafting, filed, testified, or withdrawn."),
            ("Assigned to", "Optional. The internal owner shepherding this testimony."),
            ("Notes", "Optional."),
        ],
    )
    headers = ["Docket", "Name", "Firm (optional)", "Topic (optional)",
               "Side (utility / staff / intervenor)", "Testimony due - YYYY-MM-DD",
               "Status (retained / drafting / filed / testified / withdrawn)",
               "Assigned to (optional)", "Notes (optional)"]
    _data_sheet(wb, headers,
                ["26-0347-EL-AIR  <-- EXAMPLE: overwrite or delete this row", "Dr. Jane Smith",
                 "Example Consulting", "Cost of capital", "utility", "2026-11-20",
                 "drafting", "John Doe", "EXAMPLE ROW - delete before uploading"],
                dropdowns={5: ["utility", "staff", "intervenor"],
                           7: ["retained", "drafting", "filed", "testified", "withdrawn"]})
    return _workbook_bytes(wb)


TEMPLATES = {
    "positions": ("three-way-positions-template.xlsx", positions_template,
                  "Three-way positions (as-filed vs Staff vs final)"),
    "riders": ("tariff-rider-inventory-template.xlsx", riders_template,
               "Tariff and rider inventory"),
    "cases": ("case-header-template.xlsx", cases_template, "Case header info"),
    "earned_roe": ("earned-roe-template.xlsx", earned_roe_template,
                  "Earned ROE by year (actual vs. authorized)"),
    "discovery": ("discovery-log-template.xlsx", discovery_template,
                  "Discovery log (DRs, interrogatories, depositions)"),
    "experts": ("expert-roster-template.xlsx", experts_template,
                "Expert witness roster"),
}
