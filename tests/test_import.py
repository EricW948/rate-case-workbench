"""Import validation + loading tests (positions, cases, riders) and the HTTP flow."""
from io import BytesIO

from openpyxl import Workbook, load_workbook

import database
import seed
import excel_import
import excel_templates
from models import Case, ThreeWayPosition, TariffRider, Utility


def make_wb(data_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data - fill this in"
    for r in data_rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


POS_HEADER = ["Docket", "Stage", "Period", "RR", "Change", "ROE", "Cap", "RB", "Notes"]


def test_positions_good_row_loads(db):
    seed.seed()
    wb = make_wb([POS_HEADER,
                  ["20-585-EL-AIR", "staff", "", "900", "240", "9.5", "", "", "import test"]])
    rows, errors = excel_import.validate_positions(wb, db)
    assert errors == [], errors
    assert len(rows) == 1
    assert rows[0]["revenue_change"] == 240.0
    assert rows[0]["roe"] == 9.5
    n = excel_import.load_positions(db, rows)
    assert n == 1
    # loading the same row again updates in place instead of duplicating
    n = excel_import.load_positions(db, rows)
    assert n == 1
    case = db.query(Case).filter(Case.docket == "20-585-EL-AIR").one()
    matches = [p for p in case.positions if p.stage == "staff" and not p.period]
    assert len(matches) == 1
    assert matches[0].revenue_change == 240.0


def test_positions_bad_rows_get_friendly_errors_and_load_nothing(db):
    seed.seed()
    before = db.query(ThreeWayPosition).count()
    wb = make_wb([POS_HEADER,
                  ["", "staff", "", "", "", "", "", "", ""],                 # row 2: no docket
                  ["20-585-EL-AIR", "whatever", "", "", "", "", "", "", ""],  # row 3: bad stage
                  ["20-585-EL-AIR", "as_filed", "", "", "lots", "", "", "", ""],  # row 4: bad number
                  ["99-9999-XX-XXX", "as_filed", "", "", "", "", "", "", ""],  # row 5: unknown docket
                  ["20-585-EL-AIR", "as_filed", "", "", "", "98", "", "", ""]])  # row 6: wild ROE
    rows, errors = excel_import.validate_positions(wb, db)
    assert rows == []
    assert len(errors) == 5, errors
    text = " | ".join(errors)
    assert "Row 2" in text and "Docket is missing" in text
    assert "Row 3" in text and "Stage" in text
    assert "Row 4" in text and "doesn't look like a number" in text
    assert "Row 5" in text and "can't find case" in text
    assert "Row 6" in text and "looks wrong" in text
    assert db.query(ThreeWayPosition).count() == before


def test_positions_accepts_money_formatting(db):
    seed.seed()
    wb = make_wb([POS_HEADER,
                  ["20-585-EL-AIR", "as_filed", "", "$1,200.5M", "-19", "9.84", "", "", ""]])
    rows, errors = excel_import.validate_positions(wb, db)
    assert errors == [], errors
    assert rows[0]["revenue_requirement"] == 1200.5
    assert rows[0]["revenue_change"] == -19.0


def test_positions_skips_example_and_blank_rows(db):
    seed.seed()
    wb = make_wb([POS_HEADER,
                  ["24-0468-EL-AIR  <-- EXAMPLE: overwrite or delete this row", "as_filed",
                   "", "", "", "", "", "", "EXAMPLE ROW"],
                  [],
                  ["21-0887-EL-AIR", "staff", "", "", "8.5", "", "", "", ""]])
    rows, errors = excel_import.validate_positions(wb, db)
    assert errors == [], errors
    assert len(rows) == 1
    assert rows[0]["docket"] == "21-0887-EL-AIR"


def test_riders_good_and_bad(db):
    seed.seed()
    header = ["Utility", "Name", "Type", "Status", "Revenue", "Next", "Notes"]
    wb = make_wb([header, ["AES Ohio", "Rider Test", "rider", "effective", "12.5",
                           "standalone", ""]])
    rows, errors = excel_import.validate_riders(wb, db)
    assert errors == [], errors
    assert excel_import.load_riders(db, rows) == 1
    r = db.query(TariffRider).filter(TariffRider.name == "Rider Test").one()
    assert r.disposition == "standalone" and r.annual_revenue == 12.5

    wb = make_wb([header, ["Narnia Power", "Rider X", "rider", "effective", "", "standalone", ""],
                  ["AES Ohio", "Rider Y", "rider", "effective", "", "maybe", ""]])
    rows, errors = excel_import.validate_riders(wb, db)
    assert rows == []
    assert len(errors) == 2
    assert "don't know a utility" in errors[0]
    assert "What happens next" in errors[1]


def test_cases_good_and_bad(db):
    seed.seed()
    header = ["Docket", "Utility", "Type", "Status", "Filed", "Staff", "Order",
              "TestYear", "Certain", "Notes"]
    wb = make_wb([header, ["99-0001-EL-AIR", "AES Ohio", "historical", "pending",
                           "2026-01-15", "", "", "", "", ""]])
    rows, errors = excel_import.validate_cases(wb, db)
    assert errors == [], errors
    assert excel_import.load_cases(db, rows) == 1
    c = db.query(Case).filter(Case.docket == "99-0001-EL-AIR").one()
    assert c.date_filed == "2026-01-15" and c.status == "pending"

    wb = make_wb([header, ["99-0002-EL-AIR", "AES Ohio", "future", "pending",
                           "not a date", "", "", "", "", ""]])
    rows, errors = excel_import.validate_cases(wb, db)
    assert rows == []
    assert any("Case type" in e for e in errors)
    assert any("isn't a date" in e for e in errors)


def test_not_a_workbook_gives_friendly_error(db):
    seed.seed()
    rows, errors = excel_import.validate_positions(BytesIO(b"hello"), db)
    assert rows == []
    assert errors and "Excel workbook" in errors[0]


def test_templates_open_and_have_expected_sheets():
    for key in ("positions", "riders", "cases"):
        _name, builder, _label = excel_templates.TEMPLATES[key]
        buf = builder()
        wb = load_workbook(buf)
        assert "README - start here" in wb.sheetnames
        assert "Data - fill this in" in wb.sheetnames


def test_http_flow(db):
    """End to end through the real HTTP routes: pages render, bad upload shows
    friendly errors, good upload previews then loads."""
    seed.seed()
    from fastapi.testclient import TestClient
    import main
    with TestClient(main.app) as client:
        assert client.get("/").status_code == 200
        assert "Rate Case Workbench" in client.get("/").text
        assert client.get("/cases").status_code == 200
        assert client.get("/riders").status_code == 200
        assert client.get("/charts").status_code == 200
        assert client.get("/import").status_code == 200
        r = client.get("/templates/positions")
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["content-type"]

        bad = make_wb([POS_HEADER, ["", "staff", "", "", "", "", "", "", ""]])
        r = client.post("/import/positions",
                        files={"file": ("bad.xlsx", bad.getvalue(),
                                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert r.status_code == 200
        assert "Docket is missing" in r.text
        assert "Nothing was saved" in r.text

        good = make_wb([POS_HEADER, ["20-585-EL-AIR", "as_filed", "test-period",
                                     "", "43.1", "", "", "", "http test"]])
        r = client.post("/import/positions",
                        files={"file": ("good.xlsx", good.getvalue(),
                                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert r.status_code == 200
        assert "look good" in r.text
        import re
        token = re.search(r'name="token" value="([0-9a-f]+)"', r.text).group(1)
        r = client.post("/import/positions/confirm", data={"token": token})
        assert r.status_code == 200
        assert "loaded" in r.text
        case = db.query(Case).filter(Case.docket == "20-585-EL-AIR").one()
        db.expire_all()
        assert any(p.stage == "as_filed" and p.period == "test-period"
                   and p.revenue_change == 43.1 for p in case.positions)
        assert "AES Ohio" in client.get("/").text
