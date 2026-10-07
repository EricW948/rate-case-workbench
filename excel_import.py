"""Excel import: parse, validate with plain-language errors, then load.

Nothing is loaded unless every row passes validation (all-or-nothing), and
nothing here ever raises on bad input - all problems come back as friendly
error strings naming the sheet, row, and column.
"""
import re
from datetime import datetime, date
from io import BytesIO
from openpyxl import load_workbook

DATA_SHEET = "Data - fill this in"


def _norm_enum(value, allowed):
    """Liberal enum matching: 'As Filed' -> 'as_filed'."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    key = re.sub(r"[\s\-]+", "_", str(value).strip().lower())
    for a in allowed:
        if key == a:
            return a, None
    return None, "enter one of: %s" % ", ".join(allowed)


def _norm_number(value):
    """Parse money/percent-ish input: '$48.8M', '1,200.5', '-19', '(19)'."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    if isinstance(value, (int, float)):
        return float(value), None
    s = str(value).strip().replace(",", "").replace("$", "").upper()
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1]
    s = s.rstrip("M").strip()
    try:
        num = float(s)
    except (ValueError, TypeError):
        return None, None
    return (-num if neg else num), None


def _norm_date(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d"), None
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d"), None
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m-%d-%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d"), None
        except ValueError:
            pass
    return None, "use YYYY-MM-DD, like 2026-04-01"


def _is_example_row(cells):
    return any(isinstance(c, str) and "EXAMPLE" in c.upper() for c in cells if c)


def _load_sheet(file_bytes):
    """Returns (worksheet, error). Never raises."""
    try:
        raw = file_bytes.read() if hasattr(file_bytes, "read") else file_bytes
        wb = load_workbook(BytesIO(raw), data_only=True, read_only=True)
    except Exception:
        return None, ("That file doesn't look like an Excel workbook. "
                      "Please upload an .xlsx file - the templates from the Import page are a safe bet.")
    if DATA_SHEET not in wb.sheetnames:
        return None, ("I couldn't find the '%s' sheet in that workbook. "
                      "Did you upload one of the templates from the Import page?" % DATA_SHEET)
    return wb[DATA_SHEET], None


def _rows(ws):
    """Yield (excel_row_number, [cell values]) for non-empty, non-example rows."""
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1:
            continue  # header
        cells = list(row)
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in cells):
            continue
        if _is_example_row(cells):
            continue
        yield i, cells


def _cell(cells, idx):
    return cells[idx] if idx < len(cells) else None


# ---------------------------------------------------------------- validation

STAGES = ["as_filed", "staff", "final_or_stip"]
CASE_TYPES = ["historical", "prospective"]
STATUSES = ["decided", "stipulated", "pending"]
KINDS = ["tariff", "rider"]
RIDER_STATUSES = ["effective", "proposed", "terminated"]
DISPOSITIONS = ["roll_into_base", "standalone", "undecided"]
DR_TYPES = ["data_request", "interrogatory", "production", "admission",
            "deposition", "motion", "other"]
DR_STATUSES = ["open", "in_progress", "in_review", "complete", "withdrawn"]
EXPERT_SIDES = ["utility", "staff", "intervenor"]
EXPERT_STATUSES = ["retained", "drafting", "filed", "testified", "withdrawn"]
DR_DONE = ["complete", "withdrawn"]


def validate_positions(file_bytes, db):
    """Returns (rows, errors). rows are dicts ready for load."""
    from models import Case
    ws, err = _load_sheet(file_bytes)
    if err:
        return [], [err]
    rows, errors = [], []
    for rn, cells in _rows(ws):
        docket = _cell(cells, 0)
        if not docket or not str(docket).strip():
            errors.append(f"Row {rn}: Docket is missing - every row needs a case number like 24-0468-EL-AIR.")
            continue
        docket = str(docket).strip()
        case = db.query(Case).filter(Case.docket == docket).first()
        if not case:
            errors.append(f"Row {rn}: I can't find case '{docket}' in the app. "
                          f"Add it first with the Case header template, then import positions.")
            continue
        stage, e = _norm_enum(_cell(cells, 1), STAGES)
        if e or not stage:
            errors.append(f"Row {rn}: Stage is missing or not recognized - {e or 'pick as_filed, staff, or final_or_stip'}.")
            continue
        row = {"docket": docket, "stage": stage,
               "period": (str(_cell(cells, 2)).strip() if _cell(cells, 2) else None),
               "notes": (str(_cell(cells, 8)).strip() if _cell(cells, 8) else None)}
        ok = True
        for idx, label, key, extra in [
            (3, "Revenue requirement", "revenue_requirement", "enter millions like 955.1"),
            (4, "Revenue change", "revenue_change", "enter millions like 48.8 (use a minus sign for a decrease)"),
            (5, "ROE", "roe", "enter it as a percent like 9.84"),
            (7, "Rate base", "rate_base", "enter millions like 2037.9"),
            (9, "Equity ratio", "equity_ratio", "enter it as a percent like 50.88"),
            (10, "Cost of debt", "cost_of_debt", "enter it as a percent like 5.25"),
            (11, "Operating expenses", "opex", "enter millions like 410.2"),
            (12, "Depreciation", "depreciation", "enter millions like 165.0"),
            (13, "Taxes", "taxes", "enter millions like 88.4"),
        ]:
            val, perr = _norm_number(_cell(cells, idx))
            if perr is not None or (_cell(cells, idx) not in (None, "") and val is None):
                errors.append(f"Row {rn}: {label} doesn't look like a number - {extra}. "
                              f"Leave it blank if you don't know it.")
                ok = False
            else:
                row[key] = val
        if row.get("roe") is not None and not (0 < row["roe"] < 30):
            errors.append(f"Row {rn}: ROE of {row['roe']} looks wrong - enter it as a percent like 9.84.")
            ok = False
        if row.get("equity_ratio") is not None and not (0 < row["equity_ratio"] <= 100):
            errors.append(f"Row {rn}: Equity ratio of {row['equity_ratio']} looks wrong - "
                          f"enter it as a percent like 50.88.")
            ok = False
        if row.get("cost_of_debt") is not None and not (0 <= row["cost_of_debt"] < 30):
            errors.append(f"Row {rn}: Cost of debt of {row['cost_of_debt']} looks wrong - "
                          f"enter it as a percent like 5.25.")
            ok = False
        row["capital_structure"] = (str(_cell(cells, 6)).strip() if _cell(cells, 6) else None)
        if ok:
            rows.append(row)
    return rows, errors


def validate_cases(file_bytes, db):
    from models import Utility
    ws, err = _load_sheet(file_bytes)
    if err:
        return [], [err]
    rows, errors = [], []
    for rn, cells in _rows(ws):
        docket = _cell(cells, 0)
        if not docket or not str(docket).strip():
            errors.append(f"Row {rn}: Docket is missing - every row needs a case number like 26-0347-EL-AIR.")
            continue
        docket = str(docket).strip()
        uname = _cell(cells, 1)
        if not uname or not str(uname).strip():
            errors.append(f"Row {rn}: Utility is missing - which utility does case '{docket}' belong to?")
            continue
        util = db.query(Utility).filter(Utility.name.ilike(str(uname).strip())).first()
        if not util:
            have = ", ".join(u.name for u in db.query(Utility).order_by(Utility.name).all())
            errors.append(f"Row {rn}: I don't know a utility called '{uname}'. "
                          f"Utilities in the app right now: {have}.")
            continue
        ctype, e = _norm_enum(_cell(cells, 2), CASE_TYPES)
        ok = True
        if e or not ctype:
            errors.append(f"Row {rn}: Case type is missing or not recognized - enter historical or prospective.")
            ok = False
        status, e = _norm_enum(_cell(cells, 3), STATUSES)
        if e or not status:
            errors.append(f"Row {rn}: Status is missing or not recognized - enter decided, stipulated, or pending.")
            ok = False
        row = {"docket": docket, "utility": util.name, "case_type": ctype, "status": status,
               "notes": (str(_cell(cells, 9)).strip() if _cell(cells, 9) else None)}
        for idx, label, key in [(4, "Date filed", "date_filed"), (5, "Staff report date", "date_staff_report"),
                                (6, "Order date", "date_order")]:
            val, derr = _norm_date(_cell(cells, idx))
            if derr:
                errors.append(f"Row {rn}: {label} isn't a date I recognize - {derr}. Leave it blank if unknown.")
                ok = False
            else:
                row[key] = val
        row["test_year"] = (str(_cell(cells, 7)).strip() if _cell(cells, 7) else None)
        row["date_certain"] = (str(_cell(cells, 8)).strip() if _cell(cells, 8) else None)
        if ok:
            rows.append(row)
    return rows, errors


def validate_riders(file_bytes, db):
    from models import Utility
    ws, err = _load_sheet(file_bytes)
    if err:
        return [], [err]
    rows, errors = [], []
    for rn, cells in _rows(ws):
        uname = _cell(cells, 0)
        if not uname or not str(uname).strip():
            errors.append(f"Row {rn}: Utility is missing - whose rider is this?")
            continue
        util = db.query(Utility).filter(Utility.name.ilike(str(uname).strip())).first()
        if not util:
            have = ", ".join(u.name for u in db.query(Utility).order_by(Utility.name).all())
            errors.append(f"Row {rn}: I don't know a utility called '{uname}'. "
                          f"Utilities in the app right now: {have}.")
            continue
        name = _cell(cells, 1)
        if not name or not str(name).strip():
            errors.append(f"Row {rn}: Name is missing - every tariff or rider needs a name like 'Rider DCR'.")
            continue
        kind, e = _norm_enum(_cell(cells, 2), KINDS)
        if e or not kind:
            errors.append(f"Row {rn}: Type is missing or not recognized - enter tariff or rider.")
            continue
        status, e = _norm_enum(_cell(cells, 3), RIDER_STATUSES)
        if e or not status:
            errors.append(f"Row {rn}: Status is missing or not recognized - enter effective, proposed, or terminated.")
            continue
        disp, e = _norm_enum(_cell(cells, 5), DISPOSITIONS)
        if e or not disp:
            errors.append(f"Row {rn}: 'What happens next' is missing or not recognized - enter "
                          f"roll_into_base, standalone, or undecided.")
            continue
        val, perr = _norm_number(_cell(cells, 4))
        if _cell(cells, 4) not in (None, "") and val is None:
            errors.append(f"Row {rn}: Annual revenue doesn't look like a number - enter millions like 12.5, "
                          f"or leave it blank.")
            continue
        rows.append({"utility": util.name, "name": str(name).strip(), "kind": kind,
                     "status": status, "annual_revenue": val, "disposition": disp,
                     "notes": (str(_cell(cells, 6)).strip() if _cell(cells, 6) else None)})
    return rows, errors


def validate_earned_roe(file_bytes, db):
    from models import Utility
    ws, err = _load_sheet(file_bytes)
    if err:
        return [], [err]
    rows, errors = [], []
    for rn, cells in _rows(ws):
        uname = _cell(cells, 0)
        if not uname or not str(uname).strip():
            errors.append(f"Row {rn}: Utility is missing - whose earned ROE is this?")
            continue
        util = db.query(Utility).filter(Utility.name.ilike(str(uname).strip())).first()
        if not util:
            have = ", ".join(u.name for u in db.query(Utility).order_by(Utility.name).all())
            errors.append(f"Row {rn}: I don't know a utility called '{uname}'. "
                          f"Utilities in the app right now: {have}.")
            continue
        year_raw = _cell(cells, 1)
        try:
            year = int(float(str(year_raw).strip()))
        except (ValueError, TypeError, AttributeError):
            errors.append(f"Row {rn}: Year doesn't look like a year - enter e.g. 2024.")
            continue
        if year < 1990 or year > 2100:
            errors.append(f"Row {rn}: Year {year} looks off - enter a calendar year like 2024.")
            continue
        roe, rerr = _norm_number(_cell(cells, 2))
        if _cell(cells, 2) in (None, "") or roe is None:
            errors.append(f"Row {rn}: Earned ROE is missing - enter a percent like 10.42.")
            continue
        source = str(_cell(cells, 3)).strip() if _cell(cells, 3) else None
        if not source:
            errors.append(f"Row {rn}: Source is missing - where did this number come from? "
                          f"E.g. '2024 PUCO earnings report'.")
            continue
        rows.append({"utility": util.name, "year": year, "earned_roe": roe,
                     "source": source,
                     "notes": (str(_cell(cells, 4)).strip() if _cell(cells, 4) else None)})
    return rows, errors


def validate_discovery(file_bytes, db):
    """DR / discovery log. Returns (rows, errors)."""
    from models import Case
    ws, err = _load_sheet(file_bytes)
    if err:
        return [], [err]
    rows, errors = [], []
    for rn, cells in _rows(ws):
        docket = _cell(cells, 0)
        if not docket or not str(docket).strip():
            errors.append(f"Row {rn}: Docket is missing - every row needs a case number like 26-0347-EL-AIR.")
            continue
        docket = str(docket).strip()
        case = db.query(Case).filter(Case.docket == docket).first()
        if not case:
            errors.append(f"Row {rn}: I can't find case '{docket}' in the app. "
                          f"Add it first with the Case header template, then import the discovery log.")
            continue
        rtype, e = _norm_enum(_cell(cells, 1), DR_TYPES)
        if _cell(cells, 1) not in (None, "") and e:
            errors.append(f"Row {rn}: Type '{_cell(cells, 1)}' isn't recognized - "
                          f"enter one of: {', '.join(DR_TYPES)}. Leave it blank for a plain DR.")
            continue
        party = str(_cell(cells, 2)).strip() if _cell(cells, 2) else None
        if not party:
            errors.append(f"Row {rn}: Requesting party is missing - who served this? E.g. Staff, OCC.")
            continue
        due, derr = _norm_date(_cell(cells, 6))
        if derr or not due:
            errors.append(f"Row {rn}: Due date is missing or not a date - {derr or 'use YYYY-MM-DD, like 2026-10-15'}. "
                          f"The due date drives the bottleneck view, so every item needs one.")
            continue
        served, serr = _norm_date(_cell(cells, 5))
        if serr:
            errors.append(f"Row {rn}: Served date isn't a date I recognize - use YYYY-MM-DD or leave it blank.")
            continue
        status, e = _norm_enum(_cell(cells, 8), DR_STATUSES)
        if _cell(cells, 8) not in (None, "") and e:
            errors.append(f"Row {rn}: Status '{_cell(cells, 8)}' isn't recognized - "
                          f"enter one of: {', '.join(DR_STATUSES)}. Leave it blank for open.")
            continue
        rows.append({
            "docket": docket,
            "request_type": rtype or "data_request",
            "requesting_party": party,
            "set_number": (str(_cell(cells, 3)).strip() if _cell(cells, 3) else None),
            "number": (str(_cell(cells, 4)).strip() if _cell(cells, 4) else None),
            "served_date": served,
            "due_date": due,
            "assigned_to": (str(_cell(cells, 7)).strip() if _cell(cells, 7) else None),
            "status": status or "open",
            "response_summary": (str(_cell(cells, 9)).strip() if _cell(cells, 9) else None),
            "notes": (str(_cell(cells, 10)).strip() if _cell(cells, 10) else None),
        })
    return rows, errors


def validate_experts(file_bytes, db):
    """Expert witness roster. Returns (rows, errors)."""
    from models import Case
    ws, err = _load_sheet(file_bytes)
    if err:
        return [], [err]
    rows, errors = [], []
    for rn, cells in _rows(ws):
        docket = _cell(cells, 0)
        if not docket or not str(docket).strip():
            errors.append(f"Row {rn}: Docket is missing - every row needs a case number like 26-0347-EL-AIR.")
            continue
        docket = str(docket).strip()
        case = db.query(Case).filter(Case.docket == docket).first()
        if not case:
            errors.append(f"Row {rn}: I can't find case '{docket}' in the app. "
                          f"Add it first with the Case header template, then import the roster.")
            continue
        name = str(_cell(cells, 1)).strip() if _cell(cells, 1) else None
        if not name:
            errors.append(f"Row {rn}: Name is missing - whose testimony is this?")
            continue
        side, e = _norm_enum(_cell(cells, 4), EXPERT_SIDES)
        if _cell(cells, 4) not in (None, "") and e:
            errors.append(f"Row {rn}: Side '{_cell(cells, 4)}' isn't recognized - "
                          f"enter utility, staff, or intervenor. Leave it blank for utility.")
            continue
        due, derr = _norm_date(_cell(cells, 5))
        if derr or not due:
            errors.append(f"Row {rn}: Testimony due is missing or not a date - {derr or 'use YYYY-MM-DD'}. "
                          f"The due date drives the bottleneck view, so every expert needs one.")
            continue
        status, e = _norm_enum(_cell(cells, 6), EXPERT_STATUSES)
        if _cell(cells, 6) not in (None, "") and e:
            errors.append(f"Row {rn}: Status '{_cell(cells, 6)}' isn't recognized - "
                          f"enter one of: {', '.join(EXPERT_STATUSES)}. Leave it blank for retained.")
            continue
        rows.append({
            "docket": docket,
            "name": name,
            "firm": (str(_cell(cells, 2)).strip() if _cell(cells, 2) else None),
            "topic": (str(_cell(cells, 3)).strip() if _cell(cells, 3) else None),
            "side": side or "utility",
            "testimony_due": due,
            "status": status or "retained",
            "assigned_to": (str(_cell(cells, 7)).strip() if _cell(cells, 7) else None),
            "notes": (str(_cell(cells, 8)).strip() if _cell(cells, 8) else None),
        })
    return rows, errors


# ---------------------------------------------------------------- loading

def load_positions(db, rows):
    from models import Case, ThreeWayPosition
    from seed import upsert_position
    n = 0
    for r in rows:
        case = db.query(Case).filter(Case.docket == r["docket"]).one()
        upsert_position(db, case, r["stage"], period=r.get("period"),
                        revenue_requirement=r.get("revenue_requirement"),
                        revenue_change=r.get("revenue_change"), roe=r.get("roe"),
                        capital_structure=r.get("capital_structure"),
                        equity_ratio=r.get("equity_ratio"),
                        cost_of_debt=r.get("cost_of_debt"),
                        rate_base=r.get("rate_base"), opex=r.get("opex"),
                        depreciation=r.get("depreciation"), taxes=r.get("taxes"),
                        notes=r.get("notes"))
        n += 1
    db.commit()
    return n


def load_cases(db, rows):
    from models import Utility
    from seed import get_or_create_case
    n = 0
    for r in rows:
        util = db.query(Utility).filter(Utility.name == r["utility"]).one()
        get_or_create_case(db, r["docket"], util, case_type=r["case_type"], status=r["status"],
                           date_filed=r.get("date_filed"),
                           date_staff_report=r.get("date_staff_report"),
                           date_order=r.get("date_order"), test_year=r.get("test_year"),
                           date_certain=r.get("date_certain"), notes=r.get("notes"))
        n += 1
    db.commit()
    return n


def load_riders(db, rows):
    from models import Utility
    from seed import upsert_rider
    n = 0
    for r in rows:
        util = db.query(Utility).filter(Utility.name == r["utility"]).one()
        upsert_rider(db, util, r["name"], kind=r["kind"], status=r["status"],
                     annual_revenue=r.get("annual_revenue"), disposition=r["disposition"],
                     notes=r.get("notes"))
        n += 1
    db.commit()
    return n


def load_earned_roe(db, rows):
    from models import Utility, EarnedROE
    n = 0
    for r in rows:
        util = db.query(Utility).filter(Utility.name == r["utility"]).one()
        rec = db.query(EarnedROE).filter(
            EarnedROE.utility_id == util.id, EarnedROE.year == r["year"]).first()
        if not rec:
            rec = EarnedROE(utility_id=util.id, year=r["year"])
            db.add(rec)
        rec.earned_roe = r["earned_roe"]
        rec.source = r["source"]
        rec.notes = r.get("notes")
        n += 1
    db.commit()
    return n


def load_discovery(db, rows):
    from models import Case
    from seed import upsert_discovery
    n = 0
    for r in rows:
        case = db.query(Case).filter(Case.docket == r["docket"]).one()
        upsert_discovery(db, case, r["request_type"], r["requesting_party"],
                         r["set_number"], r["number"],
                         served_date=r.get("served_date"), due_date=r.get("due_date"),
                         assigned_to=r.get("assigned_to"), status=r.get("status"),
                         response_summary=r.get("response_summary"),
                         notes=r.get("notes"))
        n += 1
    db.commit()
    return n


def load_experts(db, rows):
    from models import Case
    from seed import upsert_expert
    n = 0
    for r in rows:
        case = db.query(Case).filter(Case.docket == r["docket"]).one()
        upsert_expert(db, case, r["name"], firm=r.get("firm"), topic=r.get("topic"),
                      side=r.get("side"), testimony_due=r.get("testimony_due"),
                      status=r.get("status"), assigned_to=r.get("assigned_to"),
                      notes=r.get("notes"))
        n += 1
    db.commit()
    return n


VALIDATORS = {"positions": validate_positions, "cases": validate_cases, "riders": validate_riders,
              "earned_roe": validate_earned_roe, "discovery": validate_discovery,
              "experts": validate_experts}
LOADERS = {"positions": load_positions, "cases": load_cases, "riders": load_riders,
           "earned_roe": load_earned_roe, "discovery": load_discovery,
           "experts": load_experts}
